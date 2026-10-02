# -*- coding: utf-8 -*-
"""
cuadre_saleu_cetelem.py
-----------------------
Cuadre mensual de créditos: Sale-U (reporteLeadsconCredito*.csv/xlsx)
vs. capturas mensuales de CETELEM (JULIO_2026_FINAL.xlsx, AGOSTO..., etc.)

Lógica de cruce (por cliente):
  A) Teléfono igual + nombre parecido            -> Coincide
  B) Nombre muy parecido, teléfono distinto      -> Coincide por nombre (revisar teléfono)
  C) Teléfono igual pero nombre distinto         -> Coincide por teléfono (revisar nombre)
  D) Nombre medianamente parecido                -> Posible coincidencia (revisar)
  Se asigna primero lo más confiable (A), luego B, C y D, para que un teléfono
  mal capturado en Sale-U no "robe" el registro de otro cliente.

Salidas: Cuadre_SaleU_<fecha>.xlsx  y  Cuadre_SaleU_<fecha>.pdf

Uso en Colab: ejecuta la celda y sube todos los archivos (el reporte de Sale-U
y los Excel mensuales de Cetelem) cuando te lo pida.
"""

import io, os, re, glob, unicodedata, subprocess, sys
from datetime import datetime
from difflib import SequenceMatcher
import pandas as pd

try:
    import reportlab  # noqa
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", "reportlab"])

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak)

# =============================================================================
# CONFIGURACIÓN
# =============================================================================
PDV = "MG COLIMA"
UMBRAL_NOMBRE_CON_TEL = 0.50   # parecido mínimo de nombre si el teléfono coincide
UMBRAL_NOMBRE_FUERTE  = 0.80   # parecido para aceptar coincidencia SOLO por nombre
UMBRAL_NOMBRE_POSIBLE = 0.60   # parecido para marcar "posible coincidencia"
SIMILITUD_PALABRA     = 0.80   # dos palabras se consideran iguales (Monserrat/Montserrat)

# Cómo debería verse en Sale-U cada estatus de Cetelem
ESTATUS_ESPERADO = {
    "APROBADO": "ACEPTADO",
    "FINANCIADO": "ACEPTADO",
    "FINANCIADOS": "ACEPTADO",
    "RECHAZADO": "RECHAZADO",
    "CONTRAPROPUESTA": None,   # None = no se juzga, se marca para revisión
}

MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
         "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
         "NOVIEMBRE": 11, "DICIEMBRE": 12}

PALABRAS_VACIAS = {"DE", "DEL", "LA", "LAS", "LOS", "Y", "SA", "CV", "SAPI", "RL", "MA"}


# =============================================================================
# 1. CARGA DE ARCHIVOS
# =============================================================================
def obtener_archivos():
    """En Colab pide subir archivos; fuera de Colab usa la carpeta actual."""
    try:
        from google.colab import files  # noqa
        print("Sube el reporte de Sale-U y los Excel mensuales de Cetelem:")
        subidos = files.upload()
        rutas = list(subidos.keys())
    except ImportError:
        rutas = glob.glob("*.csv") + glob.glob("*.xlsx")
    saleu = [r for r in rutas if "lead" in r.lower() or "saleu" in r.lower()]
    cetelem = [r for r in rutas if r.lower().endswith((".xlsx", ".xls"))
               and r not in saleu and not r.startswith("Cuadre_SaleU")]
    if not saleu:
        raise FileNotFoundError("No encontré el reporte de Sale-U (nombre con 'lead').")
    if not cetelem:
        raise FileNotFoundError("No encontré capturas mensuales de Cetelem.")
    return saleu[0], cetelem


def detectar_mes(nombre_archivo, nombre_hoja):
    texto = f"{nombre_hoja} {nombre_archivo}".upper()
    mes = next((MESES[m] for m in MESES if m in texto), 0)
    anio = re.search(r"20\d{2}", texto)
    anio = int(anio.group()) if anio else datetime.now().year
    etiqueta = next((m for m in MESES if MESES[m] == mes and m != "SETIEMBRE"), "SIN MES")
    return anio * 100 + mes, f"{etiqueta[:3]} {anio}"


def _fila_encabezado(filas):
    """Busca la fila que tiene los títulos reales (la que contiene 'Nombre del Cliente')."""
    for i, fila in enumerate(filas):
        texto = quitar_acentos(" ".join(str(x) for x in fila)).upper()
        if "NOMBRE DEL CLIENTE" in texto and "TELEFONO" in texto:
            return i
    raise ValueError("No encontré la fila de encabezados ('Nombre del Cliente', 'Teléfono') en Sale-U.")


def leer_saleu(ruta):
    """Lee el reporte de Sale-U tal cual se descarga: ignora las líneas de título
    ('Leads Crédito', 'Agencia:', 'Rango:', renglones vacíos) y arranca en los encabezados."""
    if ruta.lower().endswith(".csv"):
        import csv
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                with open(ruta, encoding=enc, newline="") as f:
                    texto = f.read()
                break
            except UnicodeDecodeError:
                continue
        filas = list(csv.reader(io.StringIO(texto)))
        h = _fila_encabezado(filas)
        info = " ".join(" ".join(f) for f in filas[:h])
        df = pd.read_csv(io.StringIO(texto), dtype=str, skiprows=h, skip_blank_lines=False,
                         keep_default_na=False)
    else:
        crudo = pd.read_excel(ruta, header=None, dtype=str).fillna("")
        h = _fila_encabezado(crudo.values.tolist())
        info = " ".join(" ".join(map(str, f)) for f in crudo.values.tolist()[:h])
        df = crudo.iloc[h + 1:].copy()
        df.columns = crudo.iloc[h].tolist()
    info = re.sub(r"\s+", " ", info).strip()
    if info:
        print(f"Encabezado del reporte Sale-U: {info}")
    df = df[[c for c in df.columns if str(c).strip()]]          # columnas sin título
    df = df[df.apply(lambda r: "".join(map(str, r)).strip() != "", axis=1)]   # renglones vacíos
    df.columns = df.columns.str.strip()
    ren = {}
    for c in df.columns:
        cu = quitar_acentos(c).upper()
        if "NOMBRE DEL CLIENTE" in cu: ren[c] = "Nombre"
        elif "TELEFONO" in cu: ren[c] = "Telefono"
        elif cu == "ASESOR": ren[c] = "Asesor"
        elif "ESTATUS DEL CREDITO" in cu: ren[c] = "Estatus_SaleU"
        elif "FINANCIERA" in cu: ren[c] = "Financiera"
        elif "FECHA ALTA" in cu: ren[c] = "Fecha_Alta_Credito"
        elif "ESTATUS DE LEAD" in cu: ren[c] = "Estatus_Lead"
        elif "AUTO DE INTERES" in cu: ren[c] = "Auto_Interes"
    df = df.rename(columns=ren)
    for c in ["Nombre", "Telefono", "Asesor", "Estatus_SaleU", "Financiera",
              "Fecha_Alta_Credito", "Estatus_Lead", "Auto_Interes"]:
        if c not in df.columns:
            df[c] = ""
    df = df.fillna("")
    df["Tel_norm"] = df["Telefono"].map(normalizar_tel)
    df["Nombre_norm"] = df["Nombre"].map(normalizar_nombre)
    df["Fin_norm"] = df["Financiera"].map(lambda x: quitar_acentos(x).upper().strip())
    df["_fecha"] = df["Fecha_Alta_Credito"].map(convertir_fecha)
    df["Fecha_Alta_Credito"] = (df["_fecha"].dt.strftime("%d/%m/%Y %H:%M")
                                .where(df["_fecha"].notna(), df["Fecha_Alta_Credito"]))

    # Duplicados exactos (mismo teléfono y nombre): se queda el más reciente
    df = df.sort_values("_fecha", ascending=False)
    df["Veces_registrado"] = df.groupby(["Tel_norm", "Nombre_norm"])["Nombre"].transform("size")
    df = df.drop_duplicates(["Tel_norm", "Nombre_norm"]).reset_index(drop=True)
    df["id_s"] = range(len(df))
    return df


def convertir_fecha(x):
    """Acepta '2026-09-23 10:23:21' (formato actual de Sale-U) y '23/09/2026 10:23'."""
    x = str(x).strip()
    if not x:
        return pd.NaT
    if re.match(r"\d{4}-\d{2}-\d{2}", x):
        return pd.to_datetime(x, errors="coerce")
    return pd.to_datetime(x, dayfirst=True, errors="coerce")


def leer_cetelem(rutas):
    partes = []
    for ruta in rutas:
        xls = pd.ExcelFile(ruta)
        for hoja in xls.sheet_names:
            d = pd.read_excel(ruta, sheet_name=hoja, dtype=str).fillna("")
            d.columns = d.columns.str.strip()
            col = {quitar_acentos(c).upper(): c for c in d.columns}
            req = {"Folio": "FOLIO CCK", "Estatus_Cetelem": "STATUS",
                   "Nombre": "NOMBRE COMPLETO", "Telefono": "TELEFONO MOVIL",
                   "Vendedor": "NOMBRE DEL VENDEDOR"}
            if not all(v in col for v in req.values()):
                print(f"  ⚠ Hoja '{hoja}' de {ruta} no tiene el formato esperado, se omite.")
                continue
            orden, etiqueta = detectar_mes(ruta, hoja)
            sub = pd.DataFrame({k: d[col[v]] for k, v in req.items()})
            for extra, nombre in [("VEHICULO", "Vehiculo"),
                                  ("MONTO TOTAL A FINANCIAR", "Monto_Financiar")]:
                sub[nombre] = d[col[extra]] if extra in col else ""
            sub["Mes"] = etiqueta
            sub["_orden_mes"] = orden
            sub = sub[sub["Nombre"].str.strip() != ""]
            partes.append(sub)
    df = pd.concat(partes, ignore_index=True)
    df["Estatus_Cetelem"] = df["Estatus_Cetelem"].str.strip().str.upper()
    df["Nombre"] = df["Nombre"].str.split().str.join(" ")
    df["Tel_norm"] = df["Telefono"].map(normalizar_tel)
    df["Nombre_norm"] = df["Nombre"].map(normalizar_nombre)
    return df.sort_values(["_orden_mes", "Folio"]).reset_index(drop=True)


# =============================================================================
# 2. NORMALIZACIÓN Y SIMILITUD
# =============================================================================
def quitar_acentos(t):
    t = unicodedata.normalize("NFKD", str(t))
    return "".join(ch for ch in t if not unicodedata.combining(ch))


def normalizar_tel(t):
    t = str(t).strip()
    if re.search(r"[eE]\+", t):          # notación científica = número dañado
        return ""
    dig = re.sub(r"\D", "", t)
    if len(dig) > 10 and dig.startswith("52"):
        dig = dig[2:]
    if len(dig) > 10 and dig.startswith("1"):
        dig = dig[1:]
    return dig[-10:] if len(dig) >= 10 else ""


def normalizar_nombre(n):
    n = quitar_acentos(n).upper()
    if "@" in n or re.search(r"\.(COM|MX|NET)\b", n):
        return ""                        # se capturó un correo en lugar de nombre
    n = re.sub(r"[^A-ZÑ ]", " ", n)
    return " ".join(p for p in n.split() if len(p) > 1 and p not in PALABRAS_VACIAS)


def similitud_nombre(a, b):
    """Proporción de palabras del nombre más corto que aparecen en el más largo."""
    ta, tb = a.split(), b.split()
    if not ta or not tb:
        return 0.0, 0
    corto, largo = (ta, tb) if len(ta) <= len(tb) else (tb, ta)
    usados, coinc = set(), 0
    for p in corto:
        mejor, idx = 0, None
        for i, q in enumerate(largo):
            if i in usados:
                continue
            r = SequenceMatcher(None, p, q).ratio()
            if r > mejor:
                mejor, idx = r, i
        if mejor >= SIMILITUD_PALABRA:
            coinc += 1
            usados.add(idx)
    return coinc / len(corto), coinc


# =============================================================================
# 3. CONSOLIDAR CLIENTES DE CETELEM (un cliente puede aparecer varios meses)
# =============================================================================
def consolidar_clientes(cet):
    clientes = []   # cada uno: dict con tels, nombre, filas
    cet["id_c"] = -1
    for i, r in cet.iterrows():
        destino = None
        for c in clientes:
            sim, n = similitud_nombre(r["Nombre_norm"], c["nombre_norm"])
            mismo_tel = r["Tel_norm"] and r["Tel_norm"] in c["tels"]
            if (mismo_tel and sim >= UMBRAL_NOMBRE_CON_TEL) or (sim >= 0.9 and n >= 3):
                destino = c
                break
        if destino is None:
            destino = {"id": len(clientes), "tels": set(), "filas": []}
            clientes.append(destino)
        destino["nombre_norm"] = r["Nombre_norm"]
        destino["tels"].add(r["Tel_norm"])
        destino["filas"].append(i)
        cet.at[i, "id_c"] = destino["id"]

    filas = []
    for c in clientes:
        sub = cet.loc[c["filas"]]
        ult = sub.iloc[-1]
        filas.append({
            "id_c": c["id"],
            "Cliente_Cetelem": ult["Nombre"],
            "Nombre_norm": ult["Nombre_norm"],
            "Tel_Cetelem": ult["Tel_norm"] or ult["Telefono"],
            "tels": {t for t in c["tels"] if t},
            "Vendedor": ult["Vendedor"],
            "Folios": ", ".join(dict.fromkeys(sub["Folio"])),
            "Historial": " → ".join(f"{m}: {e}" for m, e in zip(sub["Mes"], sub["Estatus_Cetelem"])),
            "Ultimo_mes": ult["Mes"],
            "Estatus_final_Cetelem": ult["Estatus_Cetelem"],
            "Vehiculo": ult["Vehiculo"],
        })
    return pd.DataFrame(filas)


# =============================================================================
# 4. CRUCE
# =============================================================================
def cruzar(cli, su):
    candidatos = []
    for _, c in cli.iterrows():
        for _, s in su.iterrows():
            sim, n = similitud_nombre(c["Nombre_norm"], s["Nombre_norm"])
            tel = bool(s["Tel_norm"]) and s["Tel_norm"] in c["tels"]
            if tel and sim >= UMBRAL_NOMBRE_CON_TEL:
                nivel = 1
            elif not tel and sim >= UMBRAL_NOMBRE_FUERTE and n >= 2:
                nivel = 2
            elif tel:
                nivel = 3
            elif sim >= UMBRAL_NOMBRE_POSIBLE and n >= 3:
                nivel = 4
            else:
                continue
            candidatos.append((nivel, -sim, c["id_c"], s["id_s"], sim, tel))

    candidatos.sort()
    usados_c, usados_s, pares = set(), set(), {}
    for nivel, _, ic, is_, sim, tel in candidatos:
        if ic in usados_c or is_ in usados_s:
            continue
        usados_c.add(ic); usados_s.add(is_)
        pares[ic] = (is_, nivel, sim)

    tipo = {1: "Coincide (teléfono y nombre)",
            2: "Coincide por nombre – teléfono distinto",
            3: "Coincide por teléfono – nombre distinto",
            4: "Posible coincidencia por nombre"}
    tel_dueno = {t: r["Cliente_Cetelem"] for _, r in cli.iterrows() for t in r["tels"]}

    filas = []
    for _, c in cli.iterrows():
        base = c.drop(labels=["tels", "Nombre_norm"]).to_dict()
        if c["id_c"] in pares:
            is_, nivel, sim = pares[c["id_c"]]
            s = su.loc[su["id_s"] == is_].iloc[0]
            obs = []
            if nivel in (2, 4) and s["Tel_norm"] in tel_dueno and tel_dueno[s["Tel_norm"]] != c["Cliente_Cetelem"]:
                obs.append(f"El teléfono de Sale-U pertenece a {tel_dueno[s['Tel_norm']]} en Cetelem")
            elif nivel in (2, 4):
                obs.append(f"Teléfono Sale-U {s['Telefono'] or '(vacío)'} ≠ Cetelem {c['Tel_Cetelem']}")
            if nivel == 3:
                obs.append("Mismo teléfono pero el nombre en Sale-U es distinto: verificar y corregir nombre")
            if s["Fin_norm"] == "":
                obs.append("Financiera vacía en Sale-U")
            elif "CETELEM" not in s["Fin_norm"]:
                obs.append(f"En Sale-U dice financiera: {s['Financiera']}")
            if s["Veces_registrado"] > 1:
                obs.append(f"Registrado {s['Veces_registrado']} veces en Sale-U")
            est_cuadre, obs_est = comparar_estatus(c["Estatus_final_Cetelem"], s["Estatus_SaleU"])
            if obs_est:
                obs.append(obs_est)
            base.update({
                "Resultado": tipo[nivel],
                "Similitud_nombre": round(sim, 2),
                "Cliente_SaleU": s["Nombre"], "Tel_SaleU": s["Telefono"],
                "Asesor_SaleU": s["Asesor"], "Estatus_SaleU": s["Estatus_SaleU"],
                "Financiera_SaleU": s["Financiera"], "Fecha_alta_SaleU": s["Fecha_Alta_Credito"],
                "Cuadre_estatus": est_cuadre, "Observaciones": "; ".join(obs),
            })
        else:
            base.update({"Resultado": "NO registrado en Sale-U", "Similitud_nombre": None,
                         "Cliente_SaleU": "", "Tel_SaleU": "", "Asesor_SaleU": "",
                         "Estatus_SaleU": "", "Financiera_SaleU": "", "Fecha_alta_SaleU": "",
                         "Cuadre_estatus": "—", "Observaciones": "Dar de alta en Sale-U"})
        filas.append(base)
    cuadre = pd.DataFrame(filas)

    solo = su[~su["id_s"].isin(usados_s)].copy()
    solo["Clasificacion"] = solo["Fin_norm"].map(
        lambda f: "Otra financiera (no aplica)" if f and "CETELEM" not in f
        else "Cetelem en Sale-U sin captura en los meses cargados")
    solo["Observaciones"] = ""
    solo.loc[solo["Tel_norm"] == "", "Observaciones"] = "Teléfono inválido o dañado en Sale-U"
    solo.loc[solo["Nombre_norm"] == "", "Observaciones"] += " Nombre capturado como correo/ilegible"
    return cuadre, solo


def comparar_estatus(est_cet, est_su):
    esperado = ESTATUS_ESPERADO.get(est_cet, None)
    su = quitar_acentos(est_su).upper().strip()
    if esperado is None:
        return "Revisar", f"Cetelem en {est_cet.lower()}; Sale-U: {est_su or 'vacío'}"
    if su == esperado:
        return "OK", ""
    return "No cuadra", f"Estatus: Cetelem {est_cet} vs Sale-U {est_su or 'vacío'}"


# =============================================================================
# 5. EXCEL
# =============================================================================
COLORES = {"Coincide (teléfono y nombre)": "C6EFCE",
           "Coincide por nombre – teléfono distinto": "FFEB9C",
           "Coincide por teléfono – nombre distinto": "FFEB9C",
           "Posible coincidencia por nombre": "FCE4D6",
           "NO registrado en Sale-U": "FFC7CE",
           "No cuadra": "FFC7CE", "Revisar": "FFEB9C", "OK": "C6EFCE"}


def exportar_excel(ruta, resumen, cuadre, detalle, faltantes, revisar, solo):
    cols_cuadre = ["Resultado", "Cuadre_estatus", "Cliente_Cetelem", "Tel_Cetelem", "Vendedor",
                   "Folios", "Historial", "Estatus_final_Cetelem", "Cliente_SaleU", "Tel_SaleU",
                   "Asesor_SaleU", "Estatus_SaleU", "Financiera_SaleU", "Fecha_alta_SaleU",
                   "Similitud_nombre", "Observaciones"]
    cols_solo = ["Clasificacion", "Nombre", "Telefono", "Asesor", "Estatus_SaleU", "Financiera",
                 "Fecha_Alta_Credito", "Estatus_Lead", "Auto_Interes", "Veces_registrado", "Observaciones"]
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        resumen.to_excel(w, sheet_name="Resumen", index=False)
        cuadre[cols_cuadre].to_excel(w, sheet_name="Cuadre_por_cliente", index=False)
        detalle.to_excel(w, sheet_name="Detalle_por_mes", index=False)
        faltantes[cols_cuadre[2:8]].to_excel(w, sheet_name="Faltan_en_SaleU", index=False)
        revisar[cols_cuadre].to_excel(w, sheet_name="Revisar", index=False)
        solo[cols_solo].to_excel(w, sheet_name="Solo_en_SaleU", index=False)

    wb = load_workbook(ruta)
    borde = Side(style="thin", color="BFBFBF")
    for ws in wb.worksheets:
        for celda in ws[1]:
            celda.font = Font(name="Arial", bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F3864")
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.font = Font(name="Arial", size=10)
                celda.border = Border(top=borde, bottom=borde, left=borde, right=borde)
                if isinstance(celda.value, str) and celda.value in COLORES:
                    celda.fill = PatternFill("solid", fgColor=COLORES[celda.value])
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 55)
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


# =============================================================================
# 6. PDF
# =============================================================================
def exportar_pdf(ruta, resumen, faltantes, no_cuadra, revisar_cruce, solo_cet, meses):
    st = getSampleStyleSheet()
    chico = ParagraphStyle("chico", parent=st["Normal"], fontSize=7.5, leading=9)
    h1 = ParagraphStyle("h1", parent=st["Title"], fontSize=18, textColor=colors.HexColor("#1F3864"))
    h2 = ParagraphStyle("h2", parent=st["Heading2"], textColor=colors.HexColor("#1F3864"))
    txt = ParagraphStyle("txt", parent=st["Normal"], fontSize=9.5, leading=13)

    def tabla(df, cols, anchos, color="#1F3864"):
        if df.empty:
            return Paragraph("<i>Sin registros.</i>", txt)
        datos = [[Paragraph(f"<b>{c.replace('_', ' ')}</b>", chico) for c in cols]]
        for _, r in df.iterrows():
            datos.append([Paragraph(str(r[c]) if pd.notna(r[c]) else "", chico) for c in cols])
        t = Table(datos, colWidths=[a * cm for a in anchos], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    for df in (faltantes, no_cuadra, revisar_cruce):
        df["Cliente_Cetelem"] = df["Cliente_Cetelem"].astype(str)

    doc = SimpleDocTemplate(ruta, pagesize=landscape(letter), leftMargin=1.2 * cm,
                            rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    el = [Paragraph(f"Cuadre Sale-U vs Cetelem – {PDV}", h1),
          Paragraph(f"Meses cargados: {', '.join(meses)} · Generado: "
                    f"{datetime.now():%d/%m/%Y %H:%M}", txt), Spacer(1, 10)]
    el.append(tabla(resumen.assign(Valor=resumen["Valor"].astype(str)), ["Concepto", "Valor"], [12, 4]))
    el.append(Spacer(1, 10))
    el.append(Paragraph(
        "Cada cliente de Cetelem se busca en Sale-U primero por teléfono (validando que el nombre "
        "se parezca) y, si no aparece, por nombre aunque el teléfono sea distinto. Un cliente que "
        "aparece en varios meses se cuenta una sola vez con su último estatus.", txt))

    el += [PageBreak(), Paragraph(f"1. Créditos de Cetelem que NO están en Sale-U ({len(faltantes)})", h2),
           Paragraph("Estos clientes deben darse de alta en Sale-U con el asesor correspondiente.", txt),
           Spacer(1, 6)]
    if not faltantes.empty:
        por_vend = faltantes.groupby("Vendedor").size().sort_values(ascending=False)
        el.append(Paragraph("Por vendedor: " + " · ".join(f"<b>{v}</b>: {n}" for v, n in por_vend.items()), txt))
        el.append(Spacer(1, 6))
    el.append(tabla(faltantes.sort_values(["Vendedor", "Cliente_Cetelem"]),
                    ["Cliente_Cetelem", "Tel_Cetelem", "Vendedor", "Folios", "Historial"],
                    [6, 2.6, 5.5, 3.5, 7.6], "#C00000"))

    el += [PageBreak(), Paragraph(f"2. Estatus que no cuadra ({len(no_cuadra)})", h2),
           Paragraph("El cliente sí está en Sale-U, pero su estatus no corresponde con Cetelem.", txt),
           Spacer(1, 6),
           tabla(no_cuadra, ["Cliente_Cetelem", "Vendedor", "Historial", "Estatus_SaleU", "Observaciones"],
                 [5, 4.5, 6, 2.5, 7.2], "#C55A11")]

    el += [Spacer(1, 14), Paragraph(f"3. Coincidencias a verificar ({len(revisar_cruce)})", h2),
           Paragraph("Se emparejaron por nombre con teléfono distinto, o por teléfono con nombre distinto. "
                     "Confirma que sea la misma persona y corrige el dato en Sale-U.", txt),
           Spacer(1, 6),
           tabla(revisar_cruce, ["Resultado", "Cliente_Cetelem", "Tel_Cetelem", "Cliente_SaleU",
                                 "Tel_SaleU", "Observaciones"], [4.5, 4.8, 2.4, 4.5, 2.4, 6.6], "#BF8F00")]

    el += [PageBreak(), Paragraph(f"4. Registros Cetelem en Sale-U sin captura en los meses cargados ({len(solo_cet)})", h2),
           Paragraph("Pueden ser solicitudes de meses no incluidos en el cuadre, o registros mal "
                     "clasificados en Sale-U.", txt), Spacer(1, 6),
           tabla(solo_cet, ["Nombre", "Telefono", "Asesor", "Estatus_SaleU", "Fecha_Alta_Credito", "Observaciones"],
                 [5.5, 2.6, 5.5, 2.5, 3.2, 5.9], "#7F7F7F")]
    doc.build(el)


# =============================================================================
# 7. PROGRAMA PRINCIPAL
# =============================================================================
def procesar(ruta_su, rutas_cet, carpeta="."):
    """Hace todo el cuadre y devuelve un diccionario con tablas y rutas de salida."""
    su = leer_saleu(ruta_su)
    cet = leer_cetelem(rutas_cet)
    meses = list(dict.fromkeys(cet.sort_values("_orden_mes")["Mes"]))
    cli = consolidar_clientes(cet)
    cuadre, solo = cruzar(cli, su)

    detalle = cet.merge(cuadre[["id_c", "Resultado", "Cliente_SaleU", "Estatus_SaleU", "Cuadre_estatus"]],
                        on="id_c", how="left")
    detalle = detalle[["Mes", "Folio", "Estatus_Cetelem", "Nombre", "Telefono", "Vendedor",
                       "Vehiculo", "Monto_Financiar", "Resultado", "Cliente_SaleU",
                       "Estatus_SaleU", "Cuadre_estatus"]]

    encontrados = cuadre["Resultado"] != "NO registrado en Sale-U"
    faltantes = cuadre[~encontrados]
    revisar_cruce = cuadre[cuadre["Resultado"].isin([
        "Coincide por nombre – teléfono distinto", "Coincide por teléfono – nombre distinto",
        "Posible coincidencia por nombre"])]
    no_cuadra = cuadre[encontrados & (cuadre["Cuadre_estatus"] == "No cuadra")]
    revisar = cuadre[encontrados & ((cuadre["Cuadre_estatus"] != "OK") | cuadre.index.isin(revisar_cruce.index)
                                    | (cuadre["Observaciones"] != ""))]
    solo_cet = solo[solo["Clasificacion"].str.startswith("Cetelem")]

    resumen = pd.DataFrame([
        ("Solicitudes capturadas en Cetelem (todas las filas)", len(cet)),
        ("Clientes únicos en Cetelem", len(cli)),
        ("Registros en Sale-U (sin duplicados)", len(su)),
        ("Clientes Cetelem encontrados en Sale-U", int(encontrados.sum())),
        ("   · por teléfono y nombre", int((cuadre["Resultado"] == "Coincide (teléfono y nombre)").sum())),
        ("   · por nombre, teléfono distinto", int((cuadre["Resultado"] == "Coincide por nombre – teléfono distinto").sum())),
        ("   · por teléfono, nombre distinto", int((cuadre["Resultado"] == "Coincide por teléfono – nombre distinto").sum())),
        ("   · posibles (revisar)", int((cuadre["Resultado"] == "Posible coincidencia por nombre").sum())),
        ("Clientes Cetelem NO registrados en Sale-U", len(faltantes)),
        ("Porcentaje de cuadre", f"{encontrados.mean():.1%}" if len(cuadre) else "—"),
        ("Estatus que no cuadra", len(no_cuadra)),
        ("En contrapropuesta (revisar estatus)", int((cuadre["Cuadre_estatus"] == "Revisar").sum())),
        ("Cetelem en Sale-U sin captura en meses cargados", len(solo_cet)),
        ("Sale-U de otra financiera (no aplica)", len(solo) - len(solo_cet)),
    ], columns=["Concepto", "Valor"])

    sello = datetime.now().strftime("%Y%m%d_%H%M")
    ruta_xlsx = os.path.join(carpeta, f"Cuadre_SaleU_{sello}.xlsx")
    ruta_pdf = os.path.join(carpeta, f"Cuadre_SaleU_{sello}.pdf")
    exportar_excel(ruta_xlsx, resumen, cuadre, detalle, faltantes, revisar, solo)
    exportar_pdf(ruta_pdf, resumen, faltantes.copy(), no_cuadra.copy(), revisar_cruce.copy(), solo_cet, meses)
    return dict(resumen=resumen, cuadre=cuadre, detalle=detalle, faltantes=faltantes,
                no_cuadra=no_cuadra, revisar_cruce=revisar_cruce, revisar=revisar, solo=solo,
                solo_cet=solo_cet, meses=meses, xlsx=ruta_xlsx, pdf=ruta_pdf)


def main():
    ruta_su, rutas_cet = obtener_archivos()
    print(f"Sale-U: {ruta_su}\nCetelem: {', '.join(rutas_cet)}")
    r = procesar(ruta_su, rutas_cet)
    resumen, ruta_xlsx, ruta_pdf = r["resumen"], r["xlsx"], r["pdf"]

    print("\n" + resumen.to_string(index=False))
    print(f"\n✔ Archivos generados: {ruta_xlsx} y {ruta_pdf}")
    try:
        from google.colab import files
        files.download(ruta_xlsx)
        files.download(ruta_pdf)
    except ImportError:
        pass


if __name__ == "__main__":
    main()
