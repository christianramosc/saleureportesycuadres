# -*- coding: utf-8 -*-
"""
cuadre_seguimiento_cetelem.py
-----------------------------
Cuadre mensual: capturas de CETELEM (JULIO_2026_FINAL.xlsx, AGOSTO..., etc.)
vs. el Excel de seguimiento (Seguimiento_Rechazos_Financiera_MG.xlsx),
tomando SOLO las pestañas cuyo nombre contiene "COLIMA".

Orden de cruce (por cliente, de lo más confiable a lo menos):
  0) Mismo folio (y coincide teléfono o nombre)
  1) Mismo teléfono + nombre parecido
  2) Nombre muy parecido, teléfono distinto
  3) Mismo teléfono, nombre distinto
  4) Nombre medianamente parecido (posible)

Salidas: Cuadre_Seguimiento_<fecha>.xlsx y Cuadre_Seguimiento_<fecha>.pdf
Uso en Colab: ejecuta y sube el Excel de seguimiento + los Excel mensuales de Cetelem.
"""

import os, re, glob, unicodedata, subprocess, sys, warnings, calendar
from datetime import datetime, timedelta
from difflib import SequenceMatcher
import pandas as pd

warnings.filterwarnings("ignore")
try:
    import reportlab  # noqa
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", "reportlab"])

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

# =============================================================================
# CONFIGURACIÓN
# =============================================================================
PDV = "MG COLIMA"
FILTRO_HOJAS = "COLIMA"            # solo pestañas cuyo nombre contenga esto
PALABRA_ARCHIVO_SEGUIMIENTO = "SEGUIMIENTO"

UMBRAL_NOMBRE_CON_TEL = 0.50
UMBRAL_NOMBRE_FUERTE = 0.80
UMBRAL_NOMBRE_POSIBLE = 0.60
SIMILITUD_PALABRA = 0.80

# En qué pestaña debería estar cada estatus de Cetelem (texto que contiene el nombre de la hoja).
# None = debe estar capturado, pero no se juzga en qué pestaña.
HOJA_ESPERADA = {
    "RECHAZADO": "RECHAZO",
    "FINANCIADO": "FINALIZADO",
    "FINANCIADOS": "FINALIZADO",
    "APROBADO": None,
    "CONTRAPROPUESTA": None,
}

MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
         "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
         "NOVIEMBRE": 11, "DICIEMBRE": 12}
PALABRAS_VACIAS = {"DE", "DEL", "LA", "LAS", "LOS", "Y", "SA", "CV", "SAPI", "RL", "MA"}


# =============================================================================
# 1. UTILIDADES
# =============================================================================
def quitar_acentos(t):
    t = unicodedata.normalize("NFKD", str(t))
    return "".join(ch for ch in t if not unicodedata.combining(ch))


def normalizar_tel(t):
    t = str(t).strip()
    if t in ("", "nan", "None") or re.search(r"[eE]\+", t):
        return ""
    t = re.sub(r"\.0$", "", t)
    dig = re.sub(r"\D", "", t)
    if len(dig) > 10 and dig.startswith("52"):
        dig = dig[2:]
    return dig[-10:] if len(dig) >= 10 else ""


def normalizar_folio(f):
    f = re.sub(r"\.0$", "", str(f).strip())
    return f if f.isdigit() else ""


def normalizar_nombre(n):
    n = quitar_acentos(n).upper()
    if "@" in n:
        return ""
    n = re.sub(r"[^A-Z ]", " ", n)
    return " ".join(p for p in n.split() if len(p) > 1 and p not in PALABRAS_VACIAS)


def similitud_nombre(a, b):
    """Proporción de palabras del nombre más corto que aparecen en el más largo
    (no importa el orden: 'CARPIO GUZMAN VICTOR' = 'VICTOR CARPIO GUZMAN')."""
    ta, tb = a.split(), b.split()
    if not ta or not tb:
        return 0.0, 0
    corto, largo = (ta, tb) if len(ta) <= len(tb) else (tb, ta)
    usados, coinc = set(), 0
    for p in corto:
        mejor, idx = 0, None
        for i, q in enumerate(largo):
            if i not in usados:
                r = SequenceMatcher(None, p, q).ratio()
                if r > mejor:
                    mejor, idx = r, i
        if mejor >= SIMILITUD_PALABRA:
            coinc += 1
            usados.add(idx)
    return coinc / len(corto), coinc


def fecha_desde_folio(folio, anio):
    """En los folios CCK los primeros dígitos son el día del año (18265507 -> día 182)."""
    if len(folio) < 6:
        return None
    try:
        dia = int(folio[:-5])
        if 1 <= dia <= 366:
            return datetime(anio, 1, 1) + timedelta(days=dia - 1)
    except ValueError:
        pass
    return None


def buscar_col(columnas, *opciones):
    norm = {quitar_acentos(c).upper().strip(): c for c in columnas}
    for op in opciones:
        for k, c in norm.items():
            if k == op or k.startswith(op):
                return c
    return None


# =============================================================================
# 2. CARGA
# =============================================================================
def obtener_archivos():
    try:
        from google.colab import files
        print("Sube el Excel de seguimiento y los Excel mensuales de Cetelem:")
        rutas = list(files.upload().keys())
    except ImportError:
        rutas = glob.glob("*.xlsx")
    rutas = [r for r in rutas if r.lower().endswith((".xlsx", ".xls")) and not r.startswith("Cuadre_")]
    seg = [r for r in rutas if PALABRA_ARCHIVO_SEGUIMIENTO in quitar_acentos(r).upper()]
    cet = [r for r in rutas if r not in seg]
    if not seg:
        raise FileNotFoundError("No encontré el Excel de seguimiento (nombre con 'Seguimiento').")
    if not cet:
        raise FileNotFoundError("No encontré capturas mensuales de Cetelem.")
    return seg[0], cet


def detectar_mes(archivo, hoja):
    texto = quitar_acentos(f"{hoja} {archivo}").upper()
    mes = next((MESES[m] for m in MESES if m in texto), 0)
    anio = re.search(r"20\d{2}", texto)
    anio = int(anio.group()) if anio else datetime.now().year
    etiqueta = next((m for m in MESES if MESES[m] == mes and m != "SETIEMBRE"), "SIN MES")
    return anio, mes, f"{etiqueta[:3]} {anio}"


def leer_cetelem(rutas):
    partes = []
    for ruta in rutas:
        xls = pd.ExcelFile(ruta)
        for hoja in xls.sheet_names:
            d = pd.read_excel(ruta, sheet_name=hoja, dtype=str).fillna("")
            d.columns = d.columns.str.strip()
            c = {k: buscar_col(d.columns, v) for k, v in [
                ("Folio", "FOLIO"), ("Estatus_Cetelem", "STATUS"), ("Nombre", "NOMBRE COMPLETO"),
                ("Telefono", "TELEFONO MOVIL"), ("Vendedor", "NOMBRE DEL VENDEDOR")]}
            if None in c.values():
                print(f"  ⚠ Hoja '{hoja}' de {ruta} sin formato de captura Cetelem, se omite.")
                continue
            anio, mes, etiqueta = detectar_mes(ruta, hoja)
            sub = pd.DataFrame({k: d[v] for k, v in c.items()})
            sub["Mes"], sub["_anio"], sub["_mes"] = etiqueta, anio, mes
            partes.append(sub[sub["Nombre"].str.strip() != ""])
    df = pd.concat(partes, ignore_index=True)
    df["Estatus_Cetelem"] = df["Estatus_Cetelem"].str.strip().str.upper()
    df["Nombre"] = df["Nombre"].str.split().str.join(" ")
    df["Folio_norm"] = df["Folio"].map(normalizar_folio)
    df["Tel_norm"] = df["Telefono"].map(normalizar_tel)
    df["Nombre_norm"] = df["Nombre"].map(normalizar_nombre)
    return df.sort_values(["_anio", "_mes", "Folio"]).reset_index(drop=True)


def leer_seguimiento(ruta, anio_ref):
    xls = pd.ExcelFile(ruta)
    hojas = [h for h in xls.sheet_names if FILTRO_HOJAS in quitar_acentos(h).upper()]
    print(f"Pestañas de seguimiento usadas: {', '.join(hojas)}")
    partes = []
    for h in hojas:
        d = pd.read_excel(ruta, sheet_name=h, dtype=object)
        if d.empty:
            continue
        d.columns = [str(c).strip() for c in d.columns]
        c_nom = buscar_col(d.columns, "NOMBRE DE CLIENTE", "NOMBRE")
        c_tel = buscar_col(d.columns, "CELULAR", "TELEFONO")
        if not c_nom:
            continue
        c_fol = buscar_col(d.columns, "FOLIO")
        c_fin = buscar_col(d.columns, "FINANCIERA")
        c_ase = buscar_col(d.columns, "ASESOR DE VENTAS", "ASESOR ASIGNADO", "ASESOR")
        c_est = buscar_col(d.columns, "STATUS", "ESTATUS")
        c_fec = buscar_col(d.columns, "FECHA INGRESO")
        c_su = buscar_col(d.columns, "CLIENTE CON REGISTRO EN SALES")
        val = lambda col: d[col] if col else pd.Series([""] * len(d))
        sub = pd.DataFrame({
            "Hoja": h, "Fila_excel": d.index + 2,
            "Nombre": val(c_nom).astype(str).str.strip(), "Telefono": val(c_tel).astype(str),
            "Folio": val(c_fol).astype(str), "Financiera": val(c_fin).astype(str),
            "Asesor": val(c_ase).astype(str), "Estatus_seg": val(c_est).astype(str),
            "Fecha_ingreso": pd.to_datetime(val(c_fec), errors="coerce"),
            "Registro_SalesU": val(c_su).astype(str),
        })
        partes.append(sub[~sub["Nombre"].isin(["", "nan", "None"])])
    df = pd.concat(partes, ignore_index=True)
    for col in ["Telefono", "Folio", "Financiera", "Asesor", "Estatus_seg", "Registro_SalesU"]:
        df[col] = df[col].replace({"nan": "", "None": "", "NaT": ""}).str.strip()
    df["Folio_norm"] = df["Folio"].map(normalizar_folio)
    df["Tel_norm"] = df["Telefono"].map(normalizar_tel)
    df["Nombre_norm"] = df["Nombre"].map(normalizar_nombre)
    df["Fin_norm"] = df["Financiera"].map(lambda x: quitar_acentos(x).upper())

    # Fechas imposibles (p. ej. capturaron la fecha de nacimiento): se estiman con el folio
    df["Obs_fecha"] = ""
    malas = df["Fecha_ingreso"].isna() | (df["Fecha_ingreso"].dt.year < anio_ref - 1)
    for i in df.index[malas]:
        est = fecha_desde_folio(df.at[i, "Folio_norm"], anio_ref)
        if pd.notna(df.at[i, "Fecha_ingreso"]):
            df.at[i, "Obs_fecha"] = (f"Fecha de ingreso inválida ({df.at[i, 'Fecha_ingreso']:%d/%m/%Y})"
                                     + (", se estimó con el folio" if est else ""))
        if est:
            df.at[i, "Fecha_ingreso"] = est

    # Duplicados: misma hoja, folio, teléfono y nombre (sin importar orden/acentos)
    df["_llave"] = (df["Hoja"] + "|" + df["Folio_norm"] + "|" + df["Tel_norm"] + "|"
                    + df["Nombre_norm"].map(lambda n: " ".join(sorted(n.split()))) + "|" + df["Fin_norm"])
    df["Veces_capturado"] = df.groupby("_llave")["Nombre"].transform("size")
    df = df.drop_duplicates("_llave").reset_index(drop=True)
    df["id_s"] = range(len(df))
    return df


# =============================================================================
# 3. CLIENTES ÚNICOS DE CETELEM
# =============================================================================
def consolidar_clientes(cet):
    grupos = []
    cet["id_c"] = -1
    for i, r in cet.iterrows():
        destino = None
        for g in grupos:
            sim, n = similitud_nombre(r["Nombre_norm"], g["nombre"])
            if ((r["Tel_norm"] and r["Tel_norm"] in g["tels"] and sim >= UMBRAL_NOMBRE_CON_TEL)
                    or (sim >= 0.9 and n >= 3)):
                destino = g
                break
        if destino is None:
            destino = {"id": len(grupos), "tels": set(), "filas": []}
            grupos.append(destino)
        destino["nombre"] = r["Nombre_norm"]
        destino["tels"].add(r["Tel_norm"])
        destino["filas"].append(i)
        cet.at[i, "id_c"] = destino["id"]
    filas = []
    for g in grupos:
        sub = cet.loc[g["filas"]]
        ult = sub.iloc[-1]
        filas.append({
            "id_c": g["id"], "Cliente_Cetelem": ult["Nombre"], "Nombre_norm": ult["Nombre_norm"],
            "Tel_Cetelem": ult["Tel_norm"] or ult["Telefono"], "tels": {t for t in g["tels"] if t},
            "folios": set(sub["Folio_norm"]) - {""}, "Vendedor": ult["Vendedor"],
            "Folios": ", ".join(dict.fromkeys(sub["Folio"])),
            "Historial": " → ".join(f"{m}: {e}" for m, e in zip(sub["Mes"], sub["Estatus_Cetelem"])),
            "Estatus_final_Cetelem": ult["Estatus_Cetelem"],
        })
    return pd.DataFrame(filas)


# =============================================================================
# 4. CRUCE
# =============================================================================
TIPOS = {0: "Coincide por folio", 1: "Coincide (teléfono y nombre)",
         2: "Coincide por nombre – teléfono distinto", 3: "Coincide por teléfono – nombre distinto",
         4: "Posible coincidencia por nombre"}


def cruzar(cli, seg):
    # Solo filas de Cetelem (o sin financiera, como la pestaña de finalizados) son candidatas,
    # salvo que el folio coincida.
    candidatos = []
    for _, c in cli.iterrows():
        for _, s in seg.iterrows():
            es_cet = s["Fin_norm"] == "" or "CETELEM" in s["Fin_norm"]
            folio = bool(s["Folio_norm"]) and s["Folio_norm"] in c["folios"]
            sim, n = similitud_nombre(c["Nombre_norm"], s["Nombre_norm"])
            tel = bool(s["Tel_norm"]) and s["Tel_norm"] in c["tels"]
            if folio and (tel or sim >= UMBRAL_NOMBRE_CON_TEL):
                nivel = 0
            elif not es_cet:
                continue
            elif tel and sim >= UMBRAL_NOMBRE_CON_TEL:
                nivel = 1
            elif not tel and sim >= UMBRAL_NOMBRE_FUERTE and n >= 2:
                nivel = 2
            elif tel:
                nivel = 3
            elif sim >= UMBRAL_NOMBRE_POSIBLE and n >= 3:
                nivel = 4
            else:
                continue
            candidatos.append((nivel, -sim, c["id_c"], s["id_s"], sim, tel))
    candidatos.sort()

    usados_c, usados_s, principal = set(), set(), {}
    for nivel, _, ic, is_, sim, tel in candidatos:
        if ic in usados_c or is_ in usados_s:
            continue
        usados_c.add(ic); usados_s.add(is_)
        principal[ic] = (is_, nivel, sim)

    # Otras capturas de la misma persona (otra solicitud / otra pestaña)
    extras = {}
    for nivel, _, ic, is_, sim, tel in candidatos:
        if is_ not in usados_s and ic in principal and nivel <= 1:
            extras.setdefault(ic, []).append(is_)
            usados_s.add(is_)

    tel_dueno = {t: r["Cliente_Cetelem"] for _, r in cli.iterrows() for t in r["tels"]}
    S = seg.set_index("id_s")
    filas = []
    for _, c in cli.iterrows():
        base = {k: c[k] for k in ["Cliente_Cetelem", "Tel_Cetelem", "Vendedor", "Folios",
                                  "Historial", "Estatus_final_Cetelem"]}
        esperada = HOJA_ESPERADA.get(c["Estatus_final_Cetelem"])
        if c["id_c"] not in principal:
            base.update({"Resultado": "NO capturado en seguimiento", "Cuadre_pestana": "—",
                         "Pestanas_encontrado": "", "Cliente_seguimiento": "", "Tel_seguimiento": "",
                         "Folio_seguimiento": "", "Asesor_seguimiento": "", "Fecha_ingreso_seg": "",
                         "Similitud_nombre": None,
                         "Observaciones": "Capturar en la pestaña de "
                                          + ("rechazos" if esperada == "RECHAZO" else
                                             "créditos finalizados" if esperada == "FINALIZADO" else
                                             "seguimiento (confirmar pestaña)")})
            filas.append(base)
            continue
        is_, nivel, sim = principal[c["id_c"]]
        s = S.loc[is_]
        todas = [is_] + extras.get(c["id_c"], [])
        hojas = list(dict.fromkeys(S.loc[todas, "Hoja"]))
        obs = []
        if sim < UMBRAL_NOMBRE_CON_TEL:
            obs.append(f"Nombre en seguimiento distinto: {s['Nombre']}")
        if s["Folio_norm"] and s["Folio_norm"] not in c["folios"]:
            obs.append(f"Folio en seguimiento {s['Folio_norm']} no coincide con Cetelem")
        if nivel == 0 and s["Tel_norm"] and s["Tel_norm"] not in c["tels"]:
            obs.append(f"Teléfono en seguimiento {s['Telefono']} ≠ Cetelem {c['Tel_Cetelem']}")
        if nivel in (2, 4):
            if s["Tel_norm"] in tel_dueno and tel_dueno[s["Tel_norm"]] != c["Cliente_Cetelem"]:
                obs.append(f"El teléfono capturado pertenece a {tel_dueno[s['Tel_norm']]}")
            else:
                obs.append(f"Teléfono en seguimiento {s['Telefono'] or '(vacío)'} ≠ Cetelem {c['Tel_Cetelem']}")
        if s["Fin_norm"] and "CETELEM" not in s["Fin_norm"]:
            obs.append(f"Capturado con financiera {s['Financiera']}")
        dup = int(S.loc[todas, "Veces_capturado"].sum())
        if dup > 1:
            obs.append(f"Capturado {dup} veces")
        obs += [o for o in S.loc[todas, "Obs_fecha"] if o]

        if esperada is None:
            cuadre = "Revisar"
            obs.append(f"Cetelem en {c['Estatus_final_Cetelem'].lower()}: confirmar pestaña")
        elif any(esperada in quitar_acentos(h).upper() for h in hojas):
            cuadre = "OK"
        else:
            cuadre = "No cuadra"
            obs.append(f"Cetelem {c['Estatus_final_Cetelem']} pero está en '{', '.join(hojas)}'")

        base.update({
            "Resultado": TIPOS[nivel], "Cuadre_pestana": cuadre, "Pestanas_encontrado": ", ".join(hojas),
            "Cliente_seguimiento": s["Nombre"], "Tel_seguimiento": s["Telefono"],
            "Folio_seguimiento": s["Folio_norm"], "Asesor_seguimiento": s["Asesor"],
            "Fecha_ingreso_seg": s["Fecha_ingreso"].strftime("%d/%m/%Y") if pd.notna(s["Fecha_ingreso"]) else "",
            "Similitud_nombre": round(sim, 2), "Observaciones": "; ".join(obs)})
        filas.append(base)
    return pd.DataFrame(filas), seg[~seg["id_s"].isin(usados_s)].copy()


def clasificar_sobrantes(solo, ini, fin):
    def clas(r):
        if r["Fin_norm"] and "CETELEM" not in r["Fin_norm"]:
            return "Otra financiera (no aplica)"
        if pd.isna(r["Fecha_ingreso"]):
            return "Sin fecha (no se puede ubicar en el periodo)"
        if ini <= r["Fecha_ingreso"] <= fin:
            return "Cetelem en seguimiento pero NO en capturas del periodo"
        return "Cetelem fuera del periodo cargado"
    solo["Clasificacion"] = solo.apply(clas, axis=1)
    solo["Fecha_ingreso"] = solo["Fecha_ingreso"].dt.strftime("%d/%m/%Y").fillna("")
    solo["Observaciones"] = solo["Obs_fecha"]
    solo.loc[solo["Veces_capturado"] > 1, "Observaciones"] += " Capturado varias veces."
    return solo


# =============================================================================
# 5. EXCEL Y PDF
# =============================================================================
COLORES = {"Coincide por folio": "C6EFCE", "Coincide (teléfono y nombre)": "C6EFCE",
           "Coincide por nombre – teléfono distinto": "FFEB9C",
           "Coincide por teléfono – nombre distinto": "FFEB9C",
           "Posible coincidencia por nombre": "FCE4D6", "NO capturado en seguimiento": "FFC7CE",
           "No cuadra": "FFC7CE", "Revisar": "FFEB9C", "OK": "C6EFCE",
           "Cetelem en seguimiento pero NO en capturas del periodo": "FFC7CE"}

COLS_CUADRE = ["Resultado", "Cuadre_pestana", "Cliente_Cetelem", "Tel_Cetelem", "Vendedor", "Folios",
               "Historial", "Estatus_final_Cetelem", "Pestanas_encontrado", "Cliente_seguimiento",
               "Tel_seguimiento", "Folio_seguimiento", "Asesor_seguimiento", "Fecha_ingreso_seg",
               "Similitud_nombre", "Observaciones"]
COLS_SOLO = ["Clasificacion", "Hoja", "Fila_excel", "Fecha_ingreso", "Folio_norm", "Nombre", "Telefono",
             "Financiera", "Asesor", "Estatus_seg", "Veces_capturado", "Observaciones"]


def exportar_excel(ruta, resumen, cuadre, faltantes, revisar, solo):
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        resumen.to_excel(w, sheet_name="Resumen", index=False)
        cuadre[COLS_CUADRE].to_excel(w, sheet_name="Cuadre_por_cliente", index=False)
        faltantes[COLS_CUADRE[2:8]].to_excel(w, sheet_name="Faltan_en_seguimiento", index=False)
        revisar[COLS_CUADRE].to_excel(w, sheet_name="Revisar", index=False)
        solo[COLS_SOLO].rename(columns={"Folio_norm": "Folio"}).to_excel(
            w, sheet_name="Solo_en_seguimiento", index=False)
    wb = load_workbook(ruta)
    b = Side(style="thin", color="BFBFBF")
    for ws in wb.worksheets:
        for celda in ws[1]:
            celda.font = Font(name="Arial", bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F3864")
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.font = Font(name="Arial", size=10)
                celda.border = Border(top=b, bottom=b, left=b, right=b)
                if isinstance(celda.value, str) and celda.value in COLORES:
                    celda.fill = PatternFill("solid", fgColor=COLORES[celda.value])
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 55)
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


def exportar_pdf(ruta, resumen, faltantes, no_cuadra, verificar, solo_periodo, meses, hojas):
    st = getSampleStyleSheet()
    chico = ParagraphStyle("c", parent=st["Normal"], fontSize=7.5, leading=9)
    h1 = ParagraphStyle("h1", parent=st["Title"], fontSize=18, textColor=colors.HexColor("#1F3864"))
    h2 = ParagraphStyle("h2", parent=st["Heading2"], textColor=colors.HexColor("#1F3864"))
    txt = ParagraphStyle("t", parent=st["Normal"], fontSize=9.5, leading=13)

    def tabla(df, cols, anchos, color="#1F3864"):
        if df.empty:
            return Paragraph("<i>Sin registros.</i>", txt)
        datos = [[Paragraph(f"<b>{c.replace('_', ' ')}</b>", chico) for c in cols]]
        for _, r in df.iterrows():
            datos.append([Paragraph("" if pd.isna(r[c]) else str(r[c]), chico) for c in cols])
        t = Table(datos, colWidths=[a * cm for a in anchos], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color)),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        for j in range(len(cols)):
            datos[0][j].style = ParagraphStyle("hb", parent=chico, textColor=colors.white)
        return t

    doc = SimpleDocTemplate(ruta, pagesize=landscape(letter), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    el = [Paragraph(f"Cuadre Excel de seguimiento vs Cetelem – {PDV}", h1),
          Paragraph(f"Capturas Cetelem: {', '.join(meses)} · Pestañas revisadas: {', '.join(hojas)} · "
                    f"Generado: {datetime.now():%d/%m/%Y %H:%M}", txt), Spacer(1, 10),
          tabla(resumen.assign(Valor=resumen["Valor"].astype(str)), ["Concepto", "Valor"], [13, 4]),
          Spacer(1, 10),
          Paragraph("Cada cliente de Cetelem se busca en el seguimiento primero por folio, luego por "
                    "teléfono (validando el nombre) y al final por nombre aunque el teléfono sea distinto. "
                    "Un cliente que aparece en varios meses se cuenta una vez con su último estatus. "
                    "Rechazados deben estar en la pestaña de rechazos y financiados en la de créditos "
                    "finalizados.", txt)]

    el += [PageBreak(), Paragraph(f"1. Créditos de Cetelem que NO están en el seguimiento ({len(faltantes)})", h2)]
    if not faltantes.empty:
        por_est = faltantes.groupby("Estatus_final_Cetelem").size()
        por_vend = faltantes.groupby("Vendedor").size().sort_values(ascending=False)
        el.append(Paragraph("Por estatus: " + " · ".join(f"<b>{k}</b>: {v}" for k, v in por_est.items()), txt))
        el.append(Paragraph("Por vendedor: " + " · ".join(f"<b>{k}</b>: {v}" for k, v in por_vend.items()), txt))
        el.append(Spacer(1, 6))
    el.append(tabla(faltantes.sort_values(["Estatus_final_Cetelem", "Vendedor"]),
                    ["Cliente_Cetelem", "Tel_Cetelem", "Vendedor", "Folios", "Historial", "Observaciones"],
                    [5.2, 2.4, 4.8, 3.2, 5.6, 4.0], "#C00000"))

    el += [PageBreak(), Paragraph(f"2. Capturados en la pestaña equivocada ({len(no_cuadra)})", h2),
           tabla(no_cuadra, ["Cliente_Cetelem", "Historial", "Pestanas_encontrado", "Observaciones"],
                 [5.5, 6, 5, 8.7], "#C55A11"),
           Spacer(1, 14), Paragraph(f"3. Datos a corregir en el seguimiento ({len(verificar)})", h2),
           Paragraph("Coincidencias con nombre, teléfono, folio o fecha distintos, o capturas duplicadas.", txt),
           Spacer(1, 6),
           tabla(verificar, ["Resultado", "Cliente_Cetelem", "Tel_Cetelem", "Cliente_seguimiento",
                             "Tel_seguimiento", "Observaciones"], [4.2, 4.6, 2.3, 4.6, 2.3, 7.2], "#BF8F00")]

    el += [PageBreak(),
           Paragraph(f"4. En seguimiento (Cetelem) pero NO en las capturas del periodo ({len(solo_periodo)})", h2),
           Paragraph("Registros con financiera Cetelem y fecha de ingreso dentro de los meses cargados que no "
                     "aparecen en las capturas mensuales.", txt), Spacer(1, 6),
           tabla(solo_periodo, ["Hoja", "Fila_excel", "Fecha_ingreso", "Folio_norm", "Nombre", "Telefono",
                                "Asesor", "Observaciones"], [3.6, 1.4, 2.2, 2.2, 5, 2.4, 4.5, 4], "#7F7F7F")]
    doc.build(el)


# =============================================================================
# 6. PRINCIPAL
# =============================================================================
def procesar(ruta_seg, rutas_cet, carpeta="."):
    """Hace todo el cuadre y devuelve un diccionario con tablas y rutas de salida."""
    cet = leer_cetelem(rutas_cet)
    anio_ref = int(cet["_anio"].max())
    seg = leer_seguimiento(ruta_seg, anio_ref)
    hojas = list(dict.fromkeys(seg["Hoja"]))

    periodos = cet[["_anio", "_mes"]].drop_duplicates().sort_values(["_anio", "_mes"])
    a0, m0 = periodos.iloc[0]; a1, m1 = periodos.iloc[-1]
    ini = datetime(int(a0), int(m0), 1)
    fin = datetime(int(a1), int(m1), calendar.monthrange(int(a1), int(m1))[1], 23, 59)
    meses = list(dict.fromkeys(cet["Mes"]))

    cli = consolidar_clientes(cet)
    cuadre, solo = cruzar(cli, seg)
    solo = clasificar_sobrantes(solo, ini, fin)

    enc = cuadre["Resultado"] != "NO capturado en seguimiento"
    faltantes = cuadre[~enc]
    no_cuadra = cuadre[enc & (cuadre["Cuadre_pestana"] == "No cuadra")]
    verificar = cuadre[enc & ((cuadre["Resultado"] != "Coincide por folio") |
                              cuadre["Observaciones"].str.contains("distinto|≠|no coincide|veces|inválida|financiera"))
                       & ~cuadre["Observaciones"].str.fullmatch(r"Cetelem en \w+: confirmar pestaña")]
    revisar = cuadre[enc & ((cuadre["Cuadre_pestana"] != "OK") | cuadre.index.isin(verificar.index))]
    solo_periodo = solo[solo["Clasificacion"].str.startswith("Cetelem en seguimiento")]

    cuenta = lambda r: int((cuadre["Resultado"] == r).sum())
    resumen = pd.DataFrame([
        ("Solicitudes en capturas Cetelem (todas las filas)", len(cet)),
        ("Clientes únicos en Cetelem", len(cli)),
        ("Registros en pestañas de Colima del seguimiento (sin duplicados)", len(seg)),
        ("Clientes Cetelem encontrados en el seguimiento", int(enc.sum())),
        ("   · por folio", cuenta(TIPOS[0])),
        ("   · por teléfono y nombre", cuenta(TIPOS[1])),
        ("   · por nombre, teléfono distinto", cuenta(TIPOS[2])),
        ("   · por teléfono, nombre distinto", cuenta(TIPOS[3])),
        ("   · posibles (revisar)", cuenta(TIPOS[4])),
        ("Clientes Cetelem NO capturados en el seguimiento", len(faltantes)),
        ("Porcentaje de cuadre", f"{enc.mean():.1%}" if len(cuadre) else "—"),
        ("Capturados en pestaña equivocada", len(no_cuadra)),
        ("Datos a corregir en el seguimiento", len(verificar)),
        ("Cetelem en seguimiento pero no en capturas del periodo", len(solo_periodo)),
        ("Seguimiento fuera del periodo o sin fecha (informativo)",
         int(solo["Clasificacion"].str.contains("fuera|Sin fecha").sum())),
        ("Seguimiento de otras financieras (no aplica)", int((solo["Clasificacion"] == "Otra financiera (no aplica)").sum())),
    ], columns=["Concepto", "Valor"])

    sello = datetime.now().strftime("%Y%m%d_%H%M")
    rx = os.path.join(carpeta, f"Cuadre_Seguimiento_{sello}.xlsx")
    rp = os.path.join(carpeta, f"Cuadre_Seguimiento_{sello}.pdf")
    exportar_excel(rx, resumen, cuadre, faltantes, revisar, solo)
    exportar_pdf(rp, resumen, faltantes, no_cuadra, verificar, solo_periodo, meses, hojas)
    return dict(resumen=resumen, cuadre=cuadre, faltantes=faltantes, no_cuadra=no_cuadra,
                verificar=verificar, revisar=revisar, solo=solo, solo_periodo=solo_periodo,
                meses=meses, hojas=hojas, xlsx=rx, pdf=rp)


def main():
    ruta_seg, rutas_cet = obtener_archivos()
    print(f"Seguimiento: {ruta_seg}\nCetelem: {', '.join(rutas_cet)}")
    r = procesar(ruta_seg, rutas_cet)
    resumen, rx, rp = r["resumen"], r["xlsx"], r["pdf"]
    print("\n" + resumen.to_string(index=False))
    print(f"\n✔ Archivos generados: {rx} y {rp}")
    try:
        from google.colab import files
        files.download(rx); files.download(rp)
    except ImportError:
        pass


if __name__ == "__main__":
    main()
